"""Load the Celadon / Lazzoni / Luxus "Programmer Handoff" workbooks.

The workbooks are read in place from `CURALINA_SUPPLIER_DATA_DIR` and never
copied into the repository (STATUS dispatch item 35). Cleaning mirrors the
notebook the model was trained in, row for row, so product IDs and features
line up with training; `tests/integration/test_consultation_parity.py`
checks that against `model/parity.json` whenever the workbooks are present.

Every cleaning rule reports what it did (`CatalogueReport`) instead of
silently dropping or guessing:
- multi-valued tags are kept whole (split on `;`), with the known typo
  `Birght & Airy` fixed and counted;
- `Practical Touches` is split into lifestyle touches, `Seating: N` and bed size;
- a price is a number above 0 or a `$1,234` string, anything else (`0`,
  blanks, `"$1,895 / $1,963"`) is missing, never guessed;
- Celadon is artwork only and has no `Product Type`: its category is `Wall Art`;
- Lazzoni and Luxus have no SKUs filled in, so their ID is `supplier:row<n>`.
"""

from __future__ import annotations

import csv
import hashlib
import math
import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

from curalina_recommendation.consultation.quiz import BED_SIZES

SUPPLIERS: tuple[str, ...] = ("Celadon", "Lazzoni", "Luxus")
_EXTENSIONS: tuple[str, ...] = (".xlsx", ".xls", ".csv")

COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "sku": ("SKU", "Supplier SKU"),
    "product_name": ("Product Name",),
    "category": ("Product Type",),
    "retail_price": ("Retail Price",),
    "room_type": ("Room Type",),
    "design_style": ("Design Style",),
    "atmosphere": ("Atmosphere",),
    "practical_touches": ("Practical Touches",),
    "width_in": ("Width (in)",),
    "depth_in": ("Depth (in)",),
    "height_in": ("Height (in)",),
    "overview": ("Product Overview",),
}
TAG_FIXES: dict[str, str] = {"Birght & Airy": "Bright & Airy"}
_SEATING = re.compile(r"Seating:\s*(\d+)")


class CatalogueNotFoundError(Exception):
    """The configured folder, or any supplier workbook in it, is missing."""


@dataclass(frozen=True, slots=True)
class CatalogueProduct:
    product_id: str
    supplier: str
    sku: str | None
    name: str
    category: str
    price: float | None
    rooms: tuple[str, ...]
    styles: tuple[str, ...]
    atmospheres: tuple[str, ...]
    touches: tuple[str, ...]
    seating: int | None
    bed_size: str | None
    width_in: float | None
    depth_in: float | None
    height_in: float | None
    overview: str | None

    @property
    def display_name(self) -> str:
        return " ".join(self.name.split())


@dataclass
class CatalogueReport:
    rows_loaded: int = 0
    dropped_no_category_or_name: int = 0
    dropped_no_tags: int = 0
    tag_fixes: Counter[str] = field(default_factory=Counter)
    unparsed_prices: Counter[str] = field(default_factory=Counter)
    source_files: dict[str, str] = field(default_factory=dict)  # supplier -> name
    fingerprint: str = ""  # sha256 over the source files' bytes, in supplier order


@dataclass(frozen=True, slots=True)
class Catalogue:
    products: tuple[CatalogueProduct, ...]
    report: CatalogueReport


def find_supplier_files(folder: Path) -> dict[str, Path]:
    if not folder.is_dir():
        raise CatalogueNotFoundError(f"supplier data folder not found: {folder}")
    found: dict[str, Path] = {}
    for supplier in SUPPLIERS:
        matches = [
            p
            for p in folder.iterdir()
            if p.is_file()
            and p.name.lower().startswith(supplier.lower())
            and p.suffix.lower() in _EXTENSIONS
            and not p.name.startswith("~$")
        ]
        for ext in _EXTENSIONS:
            same = sorted(p for p in matches if p.suffix.lower() == ext)
            if len(same) > 1:
                raise CatalogueNotFoundError(
                    f"several {ext} files for {supplier} in {folder}: "
                    f"{[p.name for p in same]}"
                )
            if same:
                found[supplier] = same[0]
                break
        if supplier not in found:
            raise CatalogueNotFoundError(
                f"no workbook for {supplier} in {folder} "
                f"(expected a file starting with {supplier!r})"
            )
    return found


def _read_rows(path: Path) -> tuple[list[str], list[tuple[object, ...]]]:
    """Header plus data rows; trailing all-empty rows are trimmed, interior
    empty rows kept, so row positions match pandas' `read_excel`."""
    rows: list[tuple[object, ...]]
    if path.suffix.lower() == ".csv":
        for encoding in ("utf-8-sig", "cp1252"):
            try:
                with path.open(newline="", encoding=encoding) as fh:
                    rows = [tuple(r) for r in csv.reader(fh)]
                break
            except UnicodeDecodeError:
                continue
        else:  # pragma: no cover - both encodings failing is a corrupt file
            raise CatalogueNotFoundError(f"could not decode {path}")
    else:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook.worksheets[0]
            rows = [tuple(r) for r in sheet.iter_rows(values_only=True)]
        finally:
            workbook.close()
    while rows and all(_blank(v) for v in rows[-1]):
        rows.pop()
    if not rows:
        return [], []
    header = [
        " ".join(str(h).split()).lower() if h is not None else "" for h in rows[0]
    ]
    return header, rows[1:]


def _blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _text(value: object) -> str | None:
    if _blank(value):
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def split_tags(value: object, report: CatalogueReport) -> tuple[str, ...]:
    if not isinstance(value, str):
        return ()
    tags: list[str] = []
    for raw in value.split(";"):
        tag = " ".join(raw.split())
        if not tag:
            continue
        if tag in TAG_FIXES:
            report.tag_fixes[f"{tag} -> {TAG_FIXES[tag]}"] += 1
            tag = TAG_FIXES[tag]
        if tag not in tags:
            tags.append(tag)
    return tuple(tags)


def parse_price(value: object, report: CatalogueReport) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, int | float):
        if math.isnan(value):
            return None
        return float(value) if value > 0 else None
    if isinstance(value, str) and value.strip():
        cleaned = value.replace("$", "").replace(",", "").strip()
        try:
            number = float(cleaned)
        except ValueError:
            report.unparsed_prices[value] += 1
            return None
        return number if number > 0 else None
    return None


def parse_number(value: object) -> float | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, int | float):
        return None if math.isnan(value) else float(value)
    try:
        number = float(str(value).strip())
    except ValueError:
        return None
    return None if math.isnan(number) else number


def split_touches(
    tags: tuple[str, ...],
) -> tuple[tuple[str, ...], int | None, str | None]:
    touches: list[str] = []
    seating: int | None = None
    bed_size: str | None = None
    for tag in tags:
        match = _SEATING.fullmatch(tag)
        if match:
            seating = int(match.group(1))
        elif tag in BED_SIZES:
            bed_size = tag
        else:
            touches.append(tag)
    return tuple(touches), seating, bed_size


def _column(header: list[str], aliases: tuple[str, ...]) -> int | None:
    for alias in aliases:
        if alias.lower() in header:
            return header.index(alias.lower())
    return None


def load_catalogue(folder: Path) -> Catalogue:
    report = CatalogueReport()
    digest = hashlib.sha256()
    products: list[CatalogueProduct] = []
    for supplier, path in find_supplier_files(folder).items():
        report.source_files[supplier] = path.name
        digest.update(path.read_bytes())
        header, rows = _read_rows(path)
        columns = {
            name: _column(header, aliases) for name, aliases in COLUMN_ALIASES.items()
        }

        def cell(
            row: tuple[object, ...], name: str, cols: dict[str, int | None] = columns
        ) -> object:
            index = cols[name]
            return row[index] if index is not None and index < len(row) else None

        for position, row in enumerate(rows):
            report.rows_loaded += 1
            name = _text(cell(row, "product_name"))
            category = (
                "Wall Art" if supplier == "Celadon" else _text(cell(row, "category"))
            )
            if not name or not category:
                report.dropped_no_category_or_name += 1
                continue
            rooms = split_tags(cell(row, "room_type"), report)
            styles = split_tags(cell(row, "design_style"), report)
            atmospheres = split_tags(cell(row, "atmosphere"), report)
            if not (rooms and styles and atmospheres):
                report.dropped_no_tags += 1
                continue
            touches, seating, bed_size = split_touches(
                split_tags(cell(row, "practical_touches"), report)
            )
            sku = _text(cell(row, "sku"))
            sku = sku.strip() if sku else None
            source_row = position + 2  # spreadsheet row number; row 1 is the header
            products.append(
                CatalogueProduct(
                    product_id=f"{supplier}:{sku}"
                    if sku
                    else f"{supplier}:row{source_row}",
                    supplier=supplier,
                    sku=sku,
                    name=name,
                    category=category,
                    price=parse_price(cell(row, "retail_price"), report),
                    rooms=rooms,
                    styles=styles,
                    atmospheres=atmospheres,
                    touches=touches,
                    seating=seating,
                    bed_size=bed_size,
                    width_in=parse_number(cell(row, "width_in")),
                    depth_in=parse_number(cell(row, "depth_in")),
                    height_in=parse_number(cell(row, "height_in")),
                    overview=_text(cell(row, "overview")),
                )
            )
    report.fingerprint = digest.hexdigest()
    ids = [p.product_id for p in products]
    if len(ids) != len(set(ids)):
        duplicates = sorted(i for i, n in Counter(ids).items() if n > 1)
        raise CatalogueNotFoundError(f"duplicate product IDs: {duplicates[:5]}")
    return Catalogue(products=tuple(products), report=report)
