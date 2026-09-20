from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from curalina_variants.evaluation.corpus import (
    ALLOWLIST_COLUMN,
    ALLOWLIST_SHEET,
    candidate_product_images,
    compute_funnel,
    excluded_product_images,
    is_excluded_filename,
    list_folder_files,
    load_allowlisted_folder_names,
    sha256_file,
    swatch_candidates,
)


def _write_workbook(path: Path, rows: list[tuple[object, ...]]) -> None:
    workbook = openpyxl.Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)
    sheet = workbook.create_sheet(ALLOWLIST_SHEET)
    sheet.append(("title row",))
    sheet.append(
        ("Name", ALLOWLIST_COLUMN, "Deficiency", "Finish", "Price")
    )
    for row in rows:
        sheet.append(row)
    workbook.save(path)


def test_is_excluded_filename_matches_thumbnail_and_contamination() -> None:
    assert is_excluded_filename("CADEIRA-ALICE-LATERAL-150x150.webp")
    assert is_excluded_filename("ChatGPT-Image-Dec-1-2025-04_28_27-PM.webp")
    assert is_excluded_filename("Screenshot-2025-12-04-at-12.46.18-PM.webp")
    assert not is_excluded_filename("CA2.NU_.204-WB-1-W-scaled.webp")


def test_load_allowlisted_folder_names_dedupes_and_skips_workbook_row(
    tmp_path: Path,
) -> None:
    workbook_path = tmp_path / "Design 44.xlsx"
    _write_workbook(
        workbook_path,
        [
            ("Chair Option 1", "Chair", "", "L1", 100),
            ("Chair Option 2", "Chair", "", "L2", 100),
            ("Bench", "Bench", "", "L3", 200),
            ("blank link", None, "", "L4", 200),
            ("workbook row", "Design 44.xlsx", "", "L5", 0),
        ],
    )

    names = load_allowlisted_folder_names(workbook_path)

    assert names == ("Chair", "Bench")


def test_load_allowlisted_folder_names_rejects_too_few_rows(tmp_path: Path) -> None:
    workbook_path = tmp_path / "Design 44.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = ALLOWLIST_SHEET
    sheet.append(("only one row",))
    workbook.save(workbook_path)

    with pytest.raises(ValueError, match="fewer than 3 rows"):
        load_allowlisted_folder_names(workbook_path)


def test_load_allowlisted_folder_names_skips_short_rows(tmp_path: Path) -> None:
    workbook_path = tmp_path / "Design 44.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = ALLOWLIST_SHEET
    sheet.append(("title",))
    sheet.append(("Name", ALLOWLIST_COLUMN))
    sheet.append(("Name only",))  # shorter than the allowlist column index
    sheet.append(("Chair", "Chair"))
    workbook.save(workbook_path)

    assert load_allowlisted_folder_names(workbook_path) == ("Chair",)


def test_load_allowlisted_folder_names_rejects_missing_column(tmp_path: Path) -> None:
    workbook_path = tmp_path / "Design 44.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = ALLOWLIST_SHEET
    sheet.append(("title",))
    sheet.append(("Name", "Not The Right Column"))
    sheet.append(("Chair", "Chair"))
    workbook.save(workbook_path)

    with pytest.raises(ValueError, match="not found"):
        load_allowlisted_folder_names(workbook_path)


def test_candidate_and_excluded_product_images_partition_webp_files(
    tmp_path: Path,
) -> None:
    folder = tmp_path / "Chair"
    folder.mkdir()
    (folder / "CA2.NU_.204-WB-1-W-scaled.webp").write_bytes(b"a")
    (folder / "CADEIRA-ALICE-LATERAL-150x150.webp").write_bytes(b"b")
    (folder / "Screenshot-2025-12-04-at-12.46.18-PM.webp").write_bytes(b"c")
    (folder / "L92.png").write_bytes(b"d")
    (folder / "Design 44.xlsx").write_bytes(b"e")

    admitted = candidate_product_images(folder)
    dropped = excluded_product_images(folder)
    swatches = swatch_candidates(folder)
    all_files = list_folder_files(folder)

    assert [p.name for p in admitted] == ["CA2.NU_.204-WB-1-W-scaled.webp"]
    assert {p.name for p in dropped} == {
        "CADEIRA-ALICE-LATERAL-150x150.webp",
        "Screenshot-2025-12-04-at-12.46.18-PM.webp",
    }
    assert [p.name for p in swatches] == ["L92.png"]
    assert len(all_files) == 5


def test_list_folder_files_returns_empty_for_missing_folder(tmp_path: Path) -> None:
    assert list_folder_files(tmp_path / "does-not-exist") == ()


def test_compute_funnel_counts_every_stage(tmp_path: Path) -> None:
    atriani_root = tmp_path / "ATRIANI"
    chair = atriani_root / "Chair"
    chair.mkdir(parents=True)
    (chair / "good-WB-1-W-scaled.webp").write_bytes(b"a")
    (chair / "good-WB-2-W-150x150.webp").write_bytes(b"b")

    bed = atriani_root / "Curva King Bed"
    bed.mkdir(parents=True)
    (bed / "ChatGPT-Image-Dec-1-2025-04_28_27-PM.webp").write_bytes(b"c")

    empty_folder = atriani_root / "Zuma Side Table"
    empty_folder.mkdir(parents=True)

    funnel = compute_funnel(
        atriani_root, ("Chair", "Curva King Bed", "Zuma Side Table")
    )

    assert funnel.folders_allowlisted == 3
    assert funnel.folders_with_admitted_webp == 1
    assert funnel.webp_seen_total == 3
    assert funnel.webp_excluded_by_name == 2
    assert funnel.webp_admitted == 1


def test_sha256_file_is_deterministic(tmp_path: Path) -> None:
    path = tmp_path / "f.bin"
    path.write_bytes(b"hello world")

    assert sha256_file(path) == sha256_file(path)
    assert len(sha256_file(path)) == 64
